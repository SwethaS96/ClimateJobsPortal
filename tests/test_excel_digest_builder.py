"""Tests for services/excel_digest_builder.py.

Pure workbook-generation tests — no database, no email provider. Each
test builds a workbook from an in-memory notification list (the exact
shape `EmailDigestService.build_pending_digest()` produces) and inspects
the result with openpyxl.
"""

from __future__ import annotations

import shutil
from datetime import date, datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from services.excel_digest_builder import ExcelDigestBuilder, parse_reliable_date

GENERATED_AT = datetime(2026, 9, 2, 12, 0, tzinfo=timezone.utc)


def _notification(**overrides):
    base = {
        "id": 1,
        "organization_name": "ICAR - IARI",
        "title": "Project Associate-I (PA-I)",
        "job_type": "PROJECT_ASSOCIATE",
        "category": "Recruitment",
        "organization_state": "Delhi",
        "organization_country": "India",
        "first_seen": "2026-09-01T10:00:00+00:00",
        "last_seen": "2026-09-02T10:00:00+00:00",
        "application_deadline": None,
        "page_url": "https://iari.res.in/notice/1",
        "pdf_url": "https://iari.res.in/notice/1.pdf",
        "website_url": "https://iari.res.in/recruitment",
    }
    base.update(overrides)
    return base


@pytest.fixture()
def builder():
    return ExcelDigestBuilder()


@pytest.fixture()
def cleanup_paths():
    paths = []
    yield paths
    for path in paths:
        shutil.rmtree(path.parent, ignore_errors=True)


# ---------------------------------------------------------------------------
# Sheets / structure
# ---------------------------------------------------------------------------


def test_workbook_has_exactly_the_three_required_sheets(builder, cleanup_paths):
    path = builder.build([_notification()], GENERATED_AT)
    cleanup_paths.append(path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["New Opportunities", "Summary", "Closing Soon"]


def test_filename_matches_required_pattern(builder, cleanup_paths):
    path = builder.build([_notification()], GENERATED_AT)
    cleanup_paths.append(path)
    assert path.name == "ClimateJobsPortal_2026-09-02.xlsx"


def test_workbook_is_written_outside_the_repository(builder, cleanup_paths):
    import os

    path = builder.build([_notification()], GENERATED_AT)
    cleanup_paths.append(path)
    repo_root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    assert not str(path.resolve()).startswith(repo_root)


# ---------------------------------------------------------------------------
# New Opportunities — Excel Table, filtering, hyperlinks, formatting
# ---------------------------------------------------------------------------


def test_new_opportunities_contains_exactly_300_rows(builder, cleanup_paths):
    notifications = [_notification(id=i, title=f"Opening {i}") for i in range(300)]
    path = builder.build(notifications, GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    assert ws.max_row == 301  # header + 300 data rows


def test_new_opportunities_contains_exactly_1338_rows(builder, cleanup_paths):
    notifications = [_notification(id=i, title=f"Opening {i}") for i in range(1338)]
    path = builder.build(notifications, GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    assert ws.max_row == 1339  # header + 1338 data rows


def test_new_opportunities_is_a_real_excel_table_with_filtering(builder, cleanup_paths):
    path = builder.build([_notification(), _notification(id=2)], GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    assert "NewOpportunities" in ws.tables
    table = ws.tables["NewOpportunities"]
    assert table.autoFilter is not None  # filter dropdowns enabled
    assert table.ref == "A1:K3"


def test_new_opportunities_header_row_is_frozen(builder, cleanup_paths):
    path = builder.build([_notification()], GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    assert ws.freeze_panes == "A2"


def test_new_opportunities_urls_are_clickable_hyperlinks(builder, cleanup_paths):
    path = builder.build([_notification()], GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    header = [cell.value for cell in ws[1]]
    app_link_col = header.index("Application Link") + 1
    pdf_link_col = header.index("Official Notice / PDF Link") + 1
    source_col = header.index("Source Website") + 1

    assert ws.cell(row=2, column=app_link_col).hyperlink is not None
    assert ws.cell(row=2, column=app_link_col).hyperlink.target == "https://iari.res.in/notice/1"
    assert ws.cell(row=2, column=pdf_link_col).hyperlink is not None
    assert ws.cell(row=2, column=source_col).hyperlink is not None


def test_new_opportunities_missing_pdf_link_has_no_hyperlink(builder, cleanup_paths):
    path = builder.build([_notification(pdf_url=None)], GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    header = [cell.value for cell in ws[1]]
    pdf_link_col = header.index("Official Notice / PDF Link") + 1
    cell = ws.cell(row=2, column=pdf_link_col)
    assert cell.value is None
    assert cell.hyperlink is None


def test_new_opportunities_job_type_is_preserved_exactly(builder, cleanup_paths):
    notifications = [
        _notification(id=1, job_type="JRF"),
        _notification(id=2, job_type="UNKNOWN"),
        _notification(id=3, job_type=None),
    ]
    path = builder.build(notifications, GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    header = [cell.value for cell in ws[1]]
    col = header.index("Job Type") + 1
    values = [ws.cell(row=r, column=col).value for r in (2, 3, 4)]
    assert values == ["JRF", "UNKNOWN", None]


def test_new_opportunities_dates_use_sensible_number_format(builder, cleanup_paths):
    path = builder.build([_notification()], GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    header = [cell.value for cell in ws[1]]
    posted_col = header.index("Posted Date") + 1
    cell = ws.cell(row=2, column=posted_col)
    # openpyxl always reads date-formatted cells back as datetime.datetime
    # (Excel has no separate pure-date storage type) — what matters is the
    # calendar date matches and the display format is date-only.
    assert cell.value.date() == date(2026, 9, 1)
    assert cell.number_format == "yyyy-mm-dd"


def test_new_opportunities_unparseable_date_preserves_raw_string(builder, cleanup_paths):
    path = builder.build([_notification(application_deadline="Rolling basis, see notice")], GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    header = [cell.value for cell in ws[1]]
    col = header.index("Deadline") + 1
    assert ws.cell(row=2, column=col).value == "Rolling basis, see notice"


def test_new_opportunities_location_derived_from_organization_state_or_country(builder, cleanup_paths):
    notifications = [
        _notification(id=1, organization_state="Kerala", organization_country="India"),
        _notification(id=2, organization_state=None, organization_country="India"),
        _notification(id=3, organization_state=None, organization_country=None),
    ]
    path = builder.build(notifications, GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    header = [cell.value for cell in ws[1]]
    col = header.index("Location") + 1
    values = [ws.cell(row=r, column=col).value for r in (2, 3, 4)]
    assert values == ["Kerala", "India", None]


def test_new_opportunities_never_invents_a_summary_description_column(builder, cleanup_paths):
    """No 'summary'/'description' field exists on the notification model —
    the sheet must not fabricate one."""
    path = builder.build([_notification()], GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    header = [cell.value for cell in ws[1]]
    assert "Summary" not in header
    assert "Description" not in header


# ---------------------------------------------------------------------------
# Summary sheet
# ---------------------------------------------------------------------------


def test_summary_total_matches_notification_count(builder, cleanup_paths):
    notifications = [_notification(id=i) for i in range(37)]
    path = builder.build(notifications, GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["Summary"]
    values = {row[0]: row[1] for row in ws.iter_rows(values_only=True) if row[0]}
    assert values["Total opportunities"] == 37


def test_summary_job_type_counts_match_input(builder, cleanup_paths):
    notifications = [
        _notification(id=1, job_type="JRF"),
        _notification(id=2, job_type="JRF"),
        _notification(id=3, job_type="FACULTY"),
        _notification(id=4, job_type=None),
    ]
    path = builder.build(notifications, GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["Summary"]
    rows = list(ws.iter_rows(values_only=True))
    job_type_section = {}
    in_section = False
    for row in rows:
        if row[0] == "Job Type":
            in_section = True
            continue
        if in_section:
            if row[0] is None:
                break
            job_type_section[row[0]] = row[1]
    assert job_type_section == {"JRF": 2, "FACULTY": 1, "(unclassified)": 1}


def test_summary_organization_counts_match_input(builder, cleanup_paths):
    notifications = [
        _notification(id=1, organization_name="Org A"),
        _notification(id=2, organization_name="Org A"),
        _notification(id=3, organization_name="Org B"),
    ]
    path = builder.build(notifications, GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["Summary"]
    rows = list(ws.iter_rows(values_only=True))
    org_section = {}
    in_section = False
    for row in rows:
        if row[0] == "Organization":
            in_section = True
            continue
        if in_section:
            if row[0] is None:
                break
            org_section[row[0]] = row[1]
    assert org_section == {"Org A": 2, "Org B": 1}


def test_summary_deadline_stats_reflect_only_reliably_parseable_deadlines(builder, cleanup_paths):
    notifications = [
        _notification(id=1, application_deadline="2026-09-10"),  # within 14 days of GENERATED_AT
        _notification(id=2, application_deadline="2026-12-25"),  # parseable, not within 14 days
        _notification(id=3, application_deadline="see notice"),  # unparseable
        _notification(id=4, application_deadline=None),
    ]
    path = builder.build(notifications, GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["Summary"]
    values = {row[0]: row[1] for row in ws.iter_rows(values_only=True) if row[0]}
    assert values["Notifications with a reliably parseable deadline"] == 2
    assert values["Closing within 14 days"] == 1


# ---------------------------------------------------------------------------
# Closing Soon
# ---------------------------------------------------------------------------


def test_closing_soon_includes_only_within_14_days(builder, cleanup_paths):
    notifications = [
        _notification(id=1, title="Due in 5 days", application_deadline="2026-09-07"),
        _notification(id=2, title="Due in 20 days", application_deadline="2026-09-22"),
        _notification(id=3, title="Already overdue", application_deadline="2026-08-01"),
    ]
    path = builder.build(notifications, GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["Closing Soon"]
    titles = [row[1] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert titles == ["Due in 5 days"]


def test_closing_soon_never_guesses_a_deadline(builder, cleanup_paths):
    notifications = [
        _notification(id=1, title="No reliable deadline", application_deadline="apply before it's too late"),
        _notification(id=2, title="Blank deadline", application_deadline=""),
        _notification(id=3, title="Null deadline", application_deadline=None),
    ]
    path = builder.build(notifications, GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["Closing Soon"]
    assert ws.max_row == 1  # header only, zero data rows


def test_closing_soon_is_sorted_by_deadline_ascending(builder, cleanup_paths):
    notifications = [
        _notification(id=1, title="Later", application_deadline="2026-09-14"),
        _notification(id=2, title="Sooner", application_deadline="2026-09-04"),
        _notification(id=3, title="Middle", application_deadline="2026-09-09"),
    ]
    path = builder.build(notifications, GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["Closing Soon"]
    titles = [row[1] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert titles == ["Sooner", "Middle", "Later"]


def test_closing_soon_sheet_exists_with_headers_when_nothing_qualifies(builder, cleanup_paths):
    path = builder.build([_notification(application_deadline=None)], GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["Closing Soon"]
    assert ws.max_row == 1
    assert [c.value for c in ws[1]] == ["Organization", "Position / Job Title", "Deadline", "Application Link"]
    assert "ClosingSoon" in ws.tables


# ---------------------------------------------------------------------------
# parse_reliable_date — never guesses
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("2026-09-15", date(2026, 9, 15)),
        ("2026-09-15T10:00:00+00:00", date(2026, 9, 15)),
        ("15-09-2026", date(2026, 9, 15)),
        ("15/09/2026", date(2026, 9, 15)),
    ],
)
def test_parse_reliable_date_common_formats(raw, expected):
    assert parse_reliable_date(raw) == expected


@pytest.mark.parametrize("raw", [None, "", "  ", "as per notification", "TBD", "rolling"])
def test_parse_reliable_date_returns_none_for_unparseable_text(raw):
    assert parse_reliable_date(raw) is None


# ---------------------------------------------------------------------------
# Illegal-character sanitization — real scraped text occasionally contains
# ASCII control characters that OOXML/openpyxl rejects outright.
# ---------------------------------------------------------------------------


def test_illegal_control_characters_in_title_do_not_crash_generation(builder, cleanup_paths):
    dirty_title = "Walk-in-interview\x0bfor Young Professional-II (YP-II)\x1f notice"
    path = builder.build([_notification(title=dirty_title)], GENERATED_AT)  # must not raise
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    header = [c.value for c in ws[1]]
    col = header.index("Position / Job Title") + 1
    value = ws.cell(row=2, column=col).value
    assert "\x0b" not in value
    assert "\x1f" not in value
    assert "Young Professional-II (YP-II)" in value  # visible text otherwise preserved


def test_illegal_control_characters_in_organization_name_do_not_crash_generation(builder, cleanup_paths):
    path = builder.build([_notification(organization_name="ICAR\x0bIndian Institute of Millets Research")], GENERATED_AT)
    cleanup_paths.append(path)
    ws = load_workbook(path)["New Opportunities"]
    header = [c.value for c in ws[1]]
    col = header.index("Organization") + 1
    assert "\x0b" not in ws.cell(row=2, column=col).value
